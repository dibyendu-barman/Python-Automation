from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pathlib import Path

Path("reports").mkdir(exist_ok=True)

doc = Document()
title = doc.add_heading("Device Validation Report", 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph("Device: ESP32")
doc.add_paragraph("Execution: 2026-09-03")

table = doc.add_table(rows=1, cols=4)
table.style = "Table Grid"

for cell, text in zip(table.rows[0].cells,
    ["Test", "Measured", "Limit", "Result"]):
 cell.text = text

rows = [
    ("Voltage", "3.31 V", "3.20–3.40 V", "PASS"),
    ("Current", "0.42 A", "0.30–0.60 A", "PASS"),
    ("Temperature", "52.0 C", "<= 45 C", "FAIL"),
]

for row in rows:
    cells = table.add_row().cells
    for cell, text in zip(cells, row):
        cell.text = text

doc.save("reports/validation_table.docx")
print("Created: reports/validation_table.docx")
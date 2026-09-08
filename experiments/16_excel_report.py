from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

wb = Workbook()
ws = wb.active

ws.title = "Results"
headers = ["Device", "Test", "Result", "Value"]

ws.append(headers)

data = [
 ["ESP32","Voltage","PASS",3.31],
 ["ESP32","Current","PASS",0.42],
 ["ESP32","Temperature","FAIL",52.0],
 ["ESP32","UART","PASS",1],
 ["ESP32","Functional","FAIL",0],
]

for row in data:
    ws.append(row)

for cell in ws[1]:
    cell.font = Font(bold=True)
    cell.alignment = Alignment(horizontal="center")

for column, width in {"A":15,"B":20,"C":12,"D":14}.items():
    ws.column_dimensions[column].width = width

ws.freeze_panes = "A2"
ws.auto_filter.ref = ws.dimensions

summary = wb.create_sheet("Summary")
summary["A1"] = "TEST REPORT SUMMARY"
summary["A3"] = "Total Tests"
summary["B3"] = "=COUNTA(Results!B2:B1000)"
summary["A4"] = "Passed"
summary["B4"] = '=COUNTIF(Results!C2:C1000,"PASS")'
summary["A5"] = "Failed"
summary["B5"] = '=COUNTIF(Results!C2:C1000,"FAIL")'
summary["A6"] = "Pass Rate"
summary["B6"] = "=B4/B3"

wb.save("reports/engineering_test_report.xlsx")
print("Created: reports/engineering_test_report.xlsx")
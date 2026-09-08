from docx import Document
from pathlib import Path

Path("reports").mkdir(exist_ok=True)

doc = Document()
doc.add_heading("Engineering Test Report", 0)
doc.add_paragraph("Device: ESP32")
doc.add_paragraph("Firmware: v1.2.0")
doc.add_heading("Test Summary", level=1)

doc.add_paragraph("Voltage Test: PASS")
doc.add_paragraph("Current Test: PASS")
doc.add_paragraph("Temperature Test: FAIL")

doc.save("reports/test_report.docx")
print("Created: reports/test_report.docx")